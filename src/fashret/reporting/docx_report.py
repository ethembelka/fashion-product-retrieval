"""Small python-docx wrapper for the step-by-step progress reports under docs/reports/.

Every report follows the same skeleton: title block, "Where we started", "What the proposal
promised", "What we did", "Results", "What comes next". Numbers are passed in, never typed.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HEADER_FILL = "D9E2F3"


class Report:
    def __init__(self, title: str, subtitle: str, author: str, course: str) -> None:
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(course)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        h = self.doc.add_heading(title, level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(subtitle).italic = True
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"{author} · {date.today():%d.%m.%Y}").font.size = Pt(10)

    def h1(self, text: str) -> None:
        self.doc.add_heading(text, level=1)

    def h2(self, text: str) -> None:
        self.doc.add_heading(text, level=2)

    def para(self, text: str, bold_lead: str | None = None) -> None:
        p = self.doc.add_paragraph()
        if bold_lead:
            p.add_run(bold_lead).bold = True
        p.add_run(text)

    def bullets(self, items: Sequence[str]) -> None:
        for it in items:
            self.doc.add_paragraph(it, style="List Bullet")

    def numbered(self, items: Sequence[str]) -> None:
        for it in items:
            self.doc.add_paragraph(it, style="List Number")

    def note(self, text: str) -> None:
        """Grey, italic side note for 'why does this matter' explanations."""
        p = self.doc.add_paragraph()
        r = p.add_run(text)
        r.italic = True
        r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    def table(self, header: Sequence[str], rows: Sequence[Sequence[object]], caption: str | None = None) -> None:
        if caption:
            p = self.doc.add_paragraph()
            p.add_run(caption).bold = True
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, name in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = str(name)
            cell.paragraphs[0].runs[0].bold = True
            _shade(cell, HEADER_FILL)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = _fmt(v)
        self.doc.add_paragraph()

    def figure(self, path: str | Path, caption: str, width_in: float = 6.0) -> None:
        self.doc.add_picture(str(path), width=Inches(width_in))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        return path


def _fmt(v: object) -> str:
    if isinstance(v, float):
        return f"{v:,.3f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if isinstance(v, int):
        return f"{v:,}".replace(",", ".")
    return str(v)


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)
